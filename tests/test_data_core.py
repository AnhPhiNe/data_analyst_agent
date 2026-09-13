"""Behavior tests through the TabularDataCore interface."""

from __future__ import annotations

import stat
from pathlib import Path
from uuid import uuid4

import pytest
from openpyxl import Workbook

import tabular_analytics_agent.data.core as data_core_module
from tabular_analytics_agent.data import (
    DataCoreLimits,
    DatasetHandle,
    DatasetIntegrityError,
    QueryExecutionError,
    SheetSelectionError,
    TabularDataCore,
    UnsafeFileError,
    UnsafeQueryError,
    UnsupportedFileError,
)
from tabular_analytics_agent.domain import FieldKind


def write_csv(path: Path, content: str) -> Path:
    path.write_text(content, encoding="utf-8", newline="")
    return path


def ingest_tiny_sales(core: TabularDataCore, directory: Path) -> DatasetHandle:
    upload = write_csv(
        directory / "sales.csv",
        "order_id,region,revenue,event_date,email\n"
        "1,North,100,2026-01-01,alice@example.com\n"
        "2,South,150,2026-01-02,bob@example.com\n"
        "3,North,,2026-01-03,carol@example.com\n"
        "3,North,,2026-01-03,carol@example.com\n",
    )
    return core.ingest(upload)


def test_csv_inspect_ingest_profile_and_query(tmp_path: Path) -> None:
    core = TabularDataCore(tmp_path / "session")
    upload = write_csv(
        tmp_path / "sales.csv",
        "order_id,region,revenue,event_date,email\n"
        "1,North,100,2026-01-01,alice@example.com\n"
        "2,South,150,2026-01-02,bob@example.com\n"
        "3,North,,2026-01-03,carol@example.com\n"
        "3,North,,2026-01-03,carol@example.com\n",
    )

    inspection = core.inspect(upload)
    assert inspection.file_format == "csv"
    assert inspection.delimiter == ","
    assert inspection.sha256

    handle = core.ingest(upload)
    assert handle.source_path.is_file()
    assert not handle.source_path.stat().st_mode & stat.S_IWUSR
    assert handle.working_database_path.is_file()
    assert handle.source_path != upload

    profile = core.profile(handle)
    fields = {field.name: field for field in profile.fields}
    assert profile.row_count == 4
    assert profile.duplicate_row_count == 1
    assert fields["revenue"].kind is FieldKind.NUMERIC
    assert fields["revenue"].missing_count == 2
    assert fields["revenue"].numeric_summary is not None
    assert fields["revenue"].numeric_summary.mean == pytest.approx(125.0)
    assert fields["event_date"].kind is FieldKind.DATETIME
    assert fields["event_date"].temporal_summary is not None
    assert "email" in profile.pii_candidates

    result = core.query(
        handle,
        "SELECT region, SUM(revenue) AS revenue FROM dataset GROUP BY region ORDER BY region",
    )
    assert result.row_count == 2
    assert result.rows == (("North", 100.0), ("South", 150.0))
    assert not result.truncated


def test_query_supports_ctes_and_enforces_row_limit(tmp_path: Path) -> None:
    core = TabularDataCore(
        tmp_path / "session",
        DataCoreLimits(max_query_rows=1),
    )
    handle = ingest_tiny_sales(core, tmp_path)

    result = core.query(
        handle,
        "WITH totals AS (SELECT region, SUM(revenue) value FROM dataset GROUP BY region) "
        "SELECT * FROM totals ORDER BY region",
    )

    assert result.row_count == 1
    assert result.truncated


def test_query_inspection_reports_canonical_fields_and_wildcards(tmp_path: Path) -> None:
    core = TabularDataCore(tmp_path / "session")
    handle = ingest_tiny_sales(core, tmp_path)

    aggregate = core.inspect_query(
        handle,
        "select region, count(*) as rows from dataset group by region",
    )
    wildcard = core.inspect_query(handle, "SELECT * FROM dataset")

    assert aggregate.normalized_sql.startswith("SELECT region, COUNT(*)")
    assert aggregate.referenced_columns == ("region",)
    assert not aggregate.has_wildcard
    assert wildcard.has_wildcard


def test_query_rejects_nonpositive_timeout(tmp_path: Path) -> None:
    core = TabularDataCore(tmp_path / "session")
    handle = ingest_tiny_sales(core, tmp_path)

    with pytest.raises(ValueError, match="positive"):
        core.query(handle, "SELECT region FROM dataset", timeout_seconds=0)


@pytest.mark.parametrize(
    "sql",
    [
        "",
        "SELECT 1; SELECT 2",
        "DELETE FROM dataset",
        "SELECT * FROM another_table",
        "SELECT * FROM read_csv_auto('secrets.csv')",
        "SELECT revenue, COLUMNS('email') FROM dataset",
        "SELECT revenue, dataset FROM dataset",
        "SELECT revenue, TO_JSON(dataset) FROM dataset",
    ],
)
def test_query_rejects_unsafe_sql(tmp_path: Path, sql: str) -> None:
    core = TabularDataCore(tmp_path / "session")
    handle = ingest_tiny_sales(core, tmp_path)

    with pytest.raises(UnsafeQueryError):
        core.query(handle, sql)


def test_query_wraps_execution_errors(tmp_path: Path) -> None:
    core = TabularDataCore(tmp_path / "session")
    handle = ingest_tiny_sales(core, tmp_path)

    with pytest.raises(QueryExecutionError, match="missing_column"):
        core.query(handle, "SELECT missing_column FROM dataset")


def test_integrity_check_detects_source_tampering(tmp_path: Path) -> None:
    core = TabularDataCore(tmp_path / "session")
    handle = ingest_tiny_sales(core, tmp_path)
    handle.source_path.chmod(handle.source_path.stat().st_mode | stat.S_IWUSR)
    handle.source_path.write_text("tampered", encoding="utf-8")

    with pytest.raises(DatasetIntegrityError, match="hash"):
        core.profile(handle)


def test_integrity_check_rejects_working_database_from_another_identity(
    tmp_path: Path,
) -> None:
    core = TabularDataCore(tmp_path / "session")
    handle = ingest_tiny_sales(core, tmp_path)
    forged_dataset = handle.dataset.model_copy(update={"dataset_id": uuid4()})
    forged_source = handle.source_path.with_name(
        f"{forged_dataset.dataset_id}{handle.source_path.suffix}"
    )
    forged_source.write_bytes(handle.source_path.read_bytes())
    forged = handle.model_copy(update={"dataset": forged_dataset, "source_path": forged_source})

    with pytest.raises(DatasetIntegrityError, match="metadata"):
        core.profile(forged)


def test_profile_handles_blank_values_and_quality_warnings(tmp_path: Path) -> None:
    rows = ["segment,score,blank"]
    rows.extend("core,1,   " for _ in range(19))
    rows.append("edge,100,   ")
    upload = write_csv(tmp_path / "quality.csv", "\n".join(rows) + "\n")
    core = TabularDataCore(tmp_path / "session")

    profile = core.profile(core.ingest(upload))
    fields = {field.name: field for field in profile.fields}

    assert fields["blank"].missing_count == 20
    assert fields["blank"].unique_count == 0
    assert "near-constant field (>=95% same value)" in fields["segment"].warnings
    assert fields["score"].numeric_summary is not None
    assert fields["score"].numeric_summary.outlier_count == 1
    assert any("Tukey IQR" in warning for warning in fields["score"].warnings)


@pytest.mark.parametrize(
    "update",
    [
        {"source_path": Path("outside.csv")},
        {"working_database_path": Path("outside.duckdb")},
        {"table_name": "other_table"},
    ],
)
def test_integrity_check_rejects_forged_handles(tmp_path: Path, update: dict[str, object]) -> None:
    core = TabularDataCore(tmp_path / "session")
    handle = ingest_tiny_sales(core, tmp_path)
    forged = handle.model_copy(update=update)

    with pytest.raises(DatasetIntegrityError):
        core.profile(forged)


def test_session_refuses_a_second_working_dataset(tmp_path: Path) -> None:
    core = TabularDataCore(tmp_path / "session")
    first = write_csv(tmp_path / "first.csv", "value\n1\n")
    second = write_csv(tmp_path / "second.csv", "value\n2\n")
    core.ingest(first)

    with pytest.raises(DatasetIntegrityError, match="already has"):
        core.ingest(second)


def test_ingest_rolls_back_when_source_cannot_be_made_read_only(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    session_directory = tmp_path / "session"
    core = TabularDataCore(session_directory)
    upload = write_csv(tmp_path / "data.csv", "value\n1\n")

    def fail_read_only(_path: Path) -> None:
        raise PermissionError("read-only failed")

    with monkeypatch.context() as patcher:
        patcher.setattr(data_core_module, "_make_read_only", fail_read_only)
        with pytest.raises(PermissionError, match="read-only failed"):
            core.ingest(upload)

    assert not (session_directory / "working" / "1" / "dataset.duckdb").exists()
    assert core.ingest(upload).working_database_path.is_file()


def test_xlsx_requires_valid_sheet_selection(tmp_path: Path) -> None:
    workbook_path = tmp_path / "workbook.xlsx"
    workbook = Workbook()
    first = workbook.active
    assert first is not None
    first.title = "Sales"
    first.append(["region", "revenue"])
    first.append(["North", 100])
    notes = workbook.create_sheet("Notes")
    notes.append(["note"])
    notes.append(["demo"])
    workbook.save(workbook_path)

    core = TabularDataCore(tmp_path / "session")
    inspection = core.inspect(workbook_path)
    assert inspection.sheets == ("Sales", "Notes")

    with pytest.raises(SheetSelectionError, match="must be selected"):
        core.ingest(workbook_path)
    with pytest.raises(SheetSelectionError, match="Unknown"):
        core.ingest(workbook_path, sheet_name="Missing")

    handle = core.ingest(workbook_path, sheet_name="Sales")
    assert handle.selected_sheet == "Sales"
    assert core.profile(handle).row_count == 1


def test_single_sheet_xlsx_is_selected_automatically(tmp_path: Path) -> None:
    workbook_path = tmp_path / "single.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Data"
    sheet.append(["value"])
    sheet.append([1])
    workbook.save(workbook_path)

    core = TabularDataCore(tmp_path / "session")
    handle = core.ingest(workbook_path)

    assert handle.selected_sheet == "Data"


@pytest.mark.parametrize(
    ("name", "content", "error_type"),
    [
        ("data.txt", "value\n1\n", UnsupportedFileError),
        ("empty.csv", "", UnsafeFileError),
        ("empty-header.csv", ",value\n1,2\n", UnsafeFileError),
        ("duplicate.csv", "Value,value\n1,2\n", UnsafeFileError),
        ("fake.xlsx", "not a workbook", UnsafeFileError),
    ],
)
def test_upload_validation_rejects_bad_files(
    tmp_path: Path,
    name: str,
    content: str,
    error_type: type[Exception],
) -> None:
    upload = write_csv(tmp_path / name, content)
    core = TabularDataCore(tmp_path / "session")

    with pytest.raises(error_type):
        core.inspect(upload)


def test_upload_size_limit_and_csv_sheet_rejection(tmp_path: Path) -> None:
    upload = write_csv(tmp_path / "data.csv", "value\n123456789\n")
    limited = TabularDataCore(tmp_path / "limited", DataCoreLimits(max_file_bytes=5))
    with pytest.raises(UnsafeFileError, match="size limit"):
        limited.inspect(upload)

    core = TabularDataCore(tmp_path / "session")
    with pytest.raises(SheetSelectionError, match="do not accept"):
        core.ingest(upload, sheet_name="Data")
