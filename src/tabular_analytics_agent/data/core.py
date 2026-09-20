"""Secure ingestion, profiling, and read-only querying behind one small interface."""

from __future__ import annotations

import csv
import hashlib
import math
import os
import re
import shutil
import stat
import time
import unicodedata
import zipfile
from concurrent.futures import ThreadPoolExecutor
from concurrent.futures import TimeoutError as FutureTimeoutError
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path, PurePosixPath
from typing import Any
from uuid import UUID, uuid4

import duckdb
import pandas as pd
from charset_normalizer import from_bytes
from openpyxl import load_workbook

from tabular_analytics_agent.data.errors import (
    DatasetIntegrityError,
    QueryExecutionError,
    QueryTimeoutError,
    SheetSelectionError,
    UnsafeFileError,
    UnsupportedFileError,
)
from tabular_analytics_agent.data.models import (
    DATASET_TABLE,
    DataCoreLimits,
    DatasetHandle,
    QueryColumn,
    QueryInspection,
    QueryResult,
    UploadInspection,
    remaining_seconds,
)
from tabular_analytics_agent.data.sql_policy import analyze_read_only_sql, quote_identifier
from tabular_analytics_agent.domain import (
    IDENTIFIER_LIKE_WARNING,
    INCONSISTENT_SPELLING_WARNING,
    POSSIBLE_PII_WARNING,
    CategoryFrequency,
    DataProfile,
    DatasetIdentity,
    FieldKind,
    FieldProfile,
    NumericSummary,
    TemporalSummary,
)

_SUPPORTED_EXTENSIONS = {".csv", ".xlsx"}
_DATE_NAME = re.compile(r"(^|_)(date|datetime|timestamp|time)($|_)", re.IGNORECASE)
# Column names are matched after lowercasing and removing diacritics, so "Họ tên" and "ho_ten"
# match alike. Only generic personal-data words are listed: a bare "name" or "tên" is not PII,
# because product and branch names are common.
_PII_NAME = re.compile(
    r"(^|_)(email|e_mail|phone|mobile|address|first_name|last_name|full_name|"
    r"customer_id|user_id|employee_id|national_id|ssn|date_of_birth|birth_date|"
    r"(customer|employee|patient|student|contact|person)_name|"
    r"ho_ten|ho_va_ten|dia_chi|so_dien_thoai|dien_thoai|sdt|cccd|cmnd|can_cuoc|ngay_sinh|"
    r"ten_(khach_hang|nhan_vien|nguoi|benh_nhan|hoc_sinh|sinh_vien))($|_)",
    re.IGNORECASE,
)
# CSV has no file signature, so content that starts like a known binary format is refused.
_BINARY_SIGNATURES = {
    b"PK\x03\x04": "a ZIP archive, such as an XLSX workbook renamed to .csv",
    b"%PDF": "a PDF document",
    b"\x89PNG": "a PNG image",
    b"\xff\xd8\xff": "a JPEG image",
    b"\x1f\x8b": "a gzip archive",
    b"\xd0\xcf\x11\xe0": "a legacy Office document",
}
_EMAIL_VALUE = re.compile(r"^[^\s@]+@[^\s@]+\.[^\s@]+$")
_PHONE_VALUE = re.compile(r"^\+?[0-9][0-9 ()-]{7,}[0-9]$")
_ISO_DATE_VALUE = re.compile(r"^\d{4}-\d{2}-\d{2}(?:[T ]\d{2}:\d{2}(?::\d{2}(?:\.\d+)?)?)?$")
_MAX_SPELLING_EXAMPLES = 3
_MAX_SPELLING_EXAMPLE_CHARS = 40
_NUMERIC_TYPES = {
    "BIGINT",
    "DECIMAL",
    "DOUBLE",
    "FLOAT",
    "HUGEINT",
    "INTEGER",
    "REAL",
    "SMALLINT",
    "TINYINT",
    "UBIGINT",
    "UHUGEINT",
    "UINTEGER",
    "USMALLINT",
    "UTINYINT",
}


@dataclass(frozen=True, slots=True)
class FullDataRead:
    """The complete rows returned for an approved statistical input query.

    This is an internal data-core value rather than a presentation ``QueryResult``.  In
    particular, it never adds the display ``max_query_rows`` limit.  The raw values are kept in
    the parent process and are later sent to the statistical worker as a serializable payload.
    """

    sql: str
    columns: tuple[QueryColumn, ...]
    rows: tuple[tuple[Any, ...], ...]
    row_count: int
    dataset_row_count: int
    population_row_count: int
    dataset_id: UUID
    working_dataset_version: int
    duration_ms: int
    inspection: QueryInspection

    def as_query_result(self, max_rows: int) -> QueryResult:
        """Build a bounded preview without changing the full-read row count used for analysis."""
        visible_rows = self.rows[:max_rows]
        rows = tuple(tuple(_normalize_scalar(value) for value in row) for row in visible_rows)
        values: dict[str, Any] = {
            "query_id": uuid4(),
            "dataset_id": self.dataset_id,
            "working_dataset_version": self.working_dataset_version,
            "sql": self.sql,
            "columns": self.columns,
            "rows": rows,
            "row_count": len(rows),
            "truncated": self.row_count > max_rows,
            "duration_ms": self.duration_ms,
            "group_by_columns": tuple(
                name
                for name in self.inspection.group_by_columns
                if any(column.name == name for column in self.columns)
            ),
            "filters": self.inspection.filters,
            "filter_scopes": self.inspection.filter_scopes,
            "dataset_count_scope": self.inspection.dataset_count_scope,
        }
        # SQL provenance is being added to QueryResult by the SQL-policy owner.  Keep this
        # adapter compatible with both the pre-change model and the new optional field.
        if "inspection" in QueryResult.model_fields:
            values["inspection"] = self.inspection
        return QueryResult(**values)


class TabularDataCore:
    """Deep module for safe local operations on one session's tabular dataset.

    The caller supplies a session directory and receives immutable records. This module owns
    format detection, file copying, header normalization, DuckDB configuration, profiling,
    query validation, and result normalization.
    """

    def __init__(self, session_directory: Path, limits: DataCoreLimits | None = None) -> None:
        self._session_directory = session_directory.resolve()
        self._limits = limits or DataCoreLimits()

    @property
    def limits(self) -> DataCoreLimits:
        """Expose immutable resource limits to bounded analytical adapters."""
        return self._limits

    def _effective_timeout(self, timeout_seconds: float | None) -> float:
        return self._limits.effective_timeout(timeout_seconds)

    def inspect(self, upload_path: Path) -> UploadInspection:
        """Validate an upload and return facts required before ingestion."""
        path = upload_path.resolve()
        if not path.is_file():
            raise UnsupportedFileError("Upload path must refer to a file")

        extension = path.suffix.casefold()
        if extension not in _SUPPORTED_EXTENSIONS:
            raise UnsupportedFileError("Only CSV and XLSX uploads are supported")

        size_bytes = path.stat().st_size
        if size_bytes <= 0:
            raise UnsafeFileError("Upload cannot be empty")
        if size_bytes > self._limits.max_file_bytes:
            raise UnsafeFileError(
                f"Upload exceeds the {self._limits.max_file_bytes}-byte size limit"
            )

        sha256 = _hash_file(path)
        if extension == ".xlsx":
            sheets = self._inspect_xlsx(path)
            return UploadInspection(
                original_filename=path.name,
                file_format="xlsx",
                size_bytes=size_bytes,
                sha256=sha256,
                sheets=sheets,
            )

        encoding, delimiter, warnings = _inspect_csv(path)
        return UploadInspection(
            original_filename=path.name,
            file_format="csv",
            size_bytes=size_bytes,
            sha256=sha256,
            encoding=encoding,
            delimiter=delimiter,
            warnings=warnings,
        )

    def ingest(self, upload_path: Path, *, sheet_name: str | None = None) -> DatasetHandle:
        """Copy a Source Dataset and materialize Working Dataset version one."""
        inspection = self.inspect(upload_path)
        selected_sheet = _select_sheet(inspection.sheets, sheet_name)
        dataset_id = uuid4()
        source_directory = self._session_directory / "source"
        working_directory = self._session_directory / "working" / "1"
        source_directory.mkdir(parents=True, exist_ok=True)
        working_directory.mkdir(parents=True, exist_ok=True)

        extension = f".{inspection.file_format}"
        source_path = source_directory / f"{dataset_id}{extension}"
        working_database_path = working_directory / "dataset.duckdb"
        source_temp = source_path.with_suffix(f"{extension}.tmp")
        database_temp = working_database_path.with_suffix(".duckdb.tmp")

        if working_database_path.exists():
            raise DatasetIntegrityError("This Analysis Session already has a Working Dataset")

        try:
            shutil.copyfile(upload_path, source_temp)
            if _hash_file(source_temp) != inspection.sha256:
                raise DatasetIntegrityError("Upload changed while it was being ingested")
            os.replace(source_temp, source_path)

            frame = self._read_frame(source_path, inspection, selected_sheet)
            _validate_and_normalize_headers(frame)
            _coerce_unambiguous_dates(frame)
            self._write_working_database(
                frame,
                database_temp,
                dataset_id=dataset_id,
                source_sha256=inspection.sha256,
                working_dataset_version=1,
            )
            os.replace(database_temp, working_database_path)
            _make_read_only(source_path)
        except Exception:
            source_temp.unlink(missing_ok=True)
            database_temp.unlink(missing_ok=True)
            working_database_path.unlink(missing_ok=True)
            _unlink_read_only(source_path)
            raise

        identity = DatasetIdentity(
            dataset_id=dataset_id,
            original_filename=inspection.original_filename,
            sha256=inspection.sha256,
            size_bytes=inspection.size_bytes,
        )
        return DatasetHandle(
            dataset=identity,
            source_path=source_path,
            working_database_path=working_database_path,
            selected_sheet=selected_sheet,
        )

    def profile(
        self,
        handle: DatasetHandle,
        *,
        timeout_seconds: float | None = None,
    ) -> DataProfile:
        """Build a deterministic Data Profile for a persisted Working Dataset.

        Profiling runs all of its DuckDB work in one parent-owned connection.  The connection is
        interrupted if the deadline expires, so callers receive a typed timeout instead of a
        profile assembled from whichever field queries happened to finish first.
        """
        started = time.perf_counter()
        effective_timeout = self._effective_timeout(timeout_seconds)
        self._verify_handle(handle)
        remaining = remaining_seconds(started, effective_timeout)
        if remaining <= 0:
            raise QueryTimeoutError(
                f"Profiling exceeded the effective {effective_timeout}-second timeout"
            )
        connection = self._connect(handle.working_database_path)
        executor = ThreadPoolExecutor(max_workers=1, thread_name_prefix="duckdb-profile")

        def execute() -> DataProfile:
            row_count = int(
                _require_row(connection.execute("SELECT COUNT(*) FROM dataset").fetchone())[0]
            )
            distinct_count = int(
                _require_row(
                    connection.execute(
                        "SELECT COUNT(*) FROM (SELECT DISTINCT * FROM dataset)"
                    ).fetchone()
                )[0]
            )
            description = connection.execute("DESCRIBE dataset").fetchall()
            fields = tuple(
                self._profile_field(connection, str(name), str(data_type), row_count)
                for name, data_type, *_ in description
            )
            pii_candidates = tuple(
                field.name for field in fields if POSSIBLE_PII_WARNING in field.warnings
            )
            return DataProfile(
                profile_id=uuid4(),
                dataset=handle.dataset,
                row_count=row_count,
                fields=fields,
                duplicate_row_count=row_count - distinct_count,
                pii_candidates=pii_candidates,
                profile_version="1",
                created_at=datetime.now().astimezone(),
            )

        future = executor.submit(execute)
        try:
            return future.result(timeout=remaining)
        except FutureTimeoutError as exc:
            connection.interrupt()
            raise QueryTimeoutError(
                f"Profiling exceeded the effective {effective_timeout}-second timeout"
            ) from exc
        except duckdb.Error as exc:
            raise QueryExecutionError(str(exc)) from exc
        finally:
            executor.shutdown(wait=True, cancel_futures=True)
            connection.close()

    def inspect_query(
        self,
        handle: DatasetHandle,
        sql: str,
        *,
        timeout_seconds: float | None = None,
    ) -> QueryInspection:
        """Validate and bind one read-only query against this Working Dataset."""
        started = time.perf_counter()
        effective_timeout = self._effective_timeout(timeout_seconds)
        self._verify_handle(handle)
        analysis = analyze_read_only_sql(sql, allowed_table=handle.table_name)
        remaining = remaining_seconds(started, effective_timeout)
        if remaining <= 0:
            raise QueryTimeoutError(
                f"Query inspection exceeded the effective {effective_timeout}-second timeout"
            )
        connection = self._connect(handle.working_database_path)
        executor = ThreadPoolExecutor(max_workers=1, thread_name_prefix="duckdb-inspect")

        def execute() -> None:
            connection.execute(f"EXPLAIN {analysis.normalized_sql}")

        future = executor.submit(execute)
        try:
            future.result(timeout=remaining)
        except FutureTimeoutError as exc:
            connection.interrupt()
            raise QueryTimeoutError(
                f"Query inspection exceeded the effective {effective_timeout}-second timeout"
            ) from exc
        except duckdb.Error as exc:
            raise QueryExecutionError(str(exc)) from exc
        finally:
            executor.shutdown(wait=True, cancel_futures=True)
            connection.close()
        return QueryInspection(
            normalized_sql=analysis.normalized_sql,
            referenced_columns=analysis.referenced_columns,
            has_wildcard=analysis.has_wildcard,
            group_by_columns=analysis.group_by_columns,
            unaliased_outputs=analysis.unaliased_outputs,
            filters=analysis.filters,
            filter_scopes=analysis.filter_scopes,
            dataset_count_scope=analysis.dataset_count_scope,
            provenance_version=analysis.provenance_version,
            base_relations=analysis.base_relations,
            output_dependencies=analysis.output_dependencies,
        )

    def query(
        self,
        handle: DatasetHandle,
        sql: str,
        *,
        timeout_seconds: float | None = None,
    ) -> QueryResult:
        """Execute one validated read-only query against the session dataset."""
        started = time.perf_counter()
        effective_timeout = self._effective_timeout(timeout_seconds)
        remaining = remaining_seconds(started, effective_timeout)
        if remaining <= 0:
            raise QueryTimeoutError(
                f"Query exceeded the effective {effective_timeout}-second timeout"
            )
        inspection = self.inspect_query(
            handle,
            sql,
            timeout_seconds=remaining,
        )
        normalized_sql = inspection.normalized_sql
        remaining = remaining_seconds(started, effective_timeout)
        if remaining <= 0:
            raise QueryTimeoutError(
                f"Query exceeded the effective {effective_timeout}-second timeout"
            )
        connection = self._connect(handle.working_database_path)
        executor = ThreadPoolExecutor(max_workers=1, thread_name_prefix="duckdb-query")

        def execute() -> tuple[list[tuple[Any, ...]], tuple[QueryColumn, ...]]:
            wrapped = (
                f"SELECT * FROM ({normalized_sql}) AS verified_query "
                f"LIMIT {self._limits.max_query_rows + 1}"
            )
            cursor = connection.execute(wrapped)
            rows = cursor.fetchall()
            columns = tuple(
                QueryColumn(name=item[0], data_type=str(item[1])) for item in cursor.description
            )
            return rows, columns

        future = executor.submit(execute)
        try:
            raw_rows, columns = future.result(timeout=remaining)
        except FutureTimeoutError as exc:
            connection.interrupt()
            raise QueryTimeoutError(
                f"Query exceeded the effective {effective_timeout}-second timeout"
            ) from exc
        except duckdb.Error as exc:
            raise QueryExecutionError(str(exc)) from exc
        finally:
            executor.shutdown(wait=True, cancel_futures=True)
            connection.close()

        truncated = len(raw_rows) > self._limits.max_query_rows
        visible_rows = raw_rows[: self._limits.max_query_rows]
        rows = tuple(tuple(_normalize_scalar(value) for value in row) for row in visible_rows)
        duration_ms = max(0, round((time.perf_counter() - started) * 1000))
        return QueryResult(
            query_id=uuid4(),
            dataset_id=handle.dataset.dataset_id,
            working_dataset_version=handle.working_dataset_version,
            sql=normalized_sql,
            columns=columns,
            rows=rows,
            row_count=len(rows),
            truncated=truncated,
            duration_ms=duration_ms,
            group_by_columns=tuple(
                name
                for name in inspection.group_by_columns
                if any(column.name == name for column in columns)
            ),
            filters=inspection.filters,
            filter_scopes=inspection.filter_scopes,
            dataset_count_scope=inspection.dataset_count_scope,
            inspection=inspection,
        )

    def read_full(
        self,
        handle: DatasetHandle,
        sql: str,
        *,
        timeout_seconds: float | None = None,
    ) -> FullDataRead:
        """Read every row of an approved input query for a statistical calculation.

        Unlike :meth:`query`, this method never wraps the SQL in a presentation ``LIMIT``.  The
        parent process owns the connection and interrupts it when the deadline expires.  The
        returned rows are normalized immediately so the statistical worker receives only a
        serializable payload and never needs access to DuckDB.
        """
        started = time.perf_counter()
        effective_timeout = self._effective_timeout(timeout_seconds)
        remaining = remaining_seconds(started, effective_timeout)
        if remaining <= 0:
            raise QueryTimeoutError(
                f"Full-data read exceeded the effective {effective_timeout}-second timeout"
            )
        inspection = self.inspect_query(
            handle,
            sql,
            timeout_seconds=remaining,
        )
        remaining = remaining_seconds(started, effective_timeout)
        if remaining <= 0:
            raise QueryTimeoutError(
                f"Full-data read exceeded the effective {effective_timeout}-second timeout"
            )
        connection = self._connect(handle.working_database_path)
        executor = ThreadPoolExecutor(max_workers=1, thread_name_prefix="duckdb-full-read")

        def execute() -> tuple[int, list[tuple[Any, ...]], tuple[QueryColumn, ...]]:
            dataset_row_count = int(
                _require_row(
                    connection.execute(
                        f"SELECT COUNT(*) FROM {quote_identifier(handle.table_name)}"
                    ).fetchone()
                )[0]
            )
            cursor = connection.execute(inspection.normalized_sql)
            raw_rows = cursor.fetchall()
            columns = tuple(
                QueryColumn(name=item[0], data_type=str(item[1])) for item in cursor.description
            )
            return dataset_row_count, raw_rows, columns

        future = executor.submit(execute)
        try:
            dataset_row_count, raw_rows, columns = future.result(timeout=remaining)
        except FutureTimeoutError as exc:
            connection.interrupt()
            raise QueryTimeoutError(
                f"Full-data read exceeded the effective {effective_timeout}-second timeout"
            ) from exc
        except MemoryError as exc:
            connection.interrupt()
            raise QueryExecutionError(
                "Full-data statistical input exceeded available memory"
            ) from exc
        except duckdb.Error as exc:
            raise QueryExecutionError(str(exc)) from exc
        finally:
            executor.shutdown(wait=True, cancel_futures=True)
            connection.close()

        # Conversion belongs to the same end-to-end deadline, but runs in the parent thread so a
        # timeout cannot leave a Python normalization loop alive in an abandoned executor thread.
        try:
            rows = _normalize_rows_with_deadline(
                raw_rows,
                deadline=started + effective_timeout,
            )
        except QueryTimeoutError:
            raise
        except MemoryError as exc:
            raise QueryExecutionError(
                "Full-data statistical input exceeded available memory"
            ) from exc

        duration_ms = max(0, round((time.perf_counter() - started) * 1000))
        return FullDataRead(
            sql=inspection.normalized_sql,
            columns=columns,
            rows=rows,
            row_count=len(rows),
            dataset_row_count=dataset_row_count,
            population_row_count=len(rows),
            dataset_id=handle.dataset.dataset_id,
            working_dataset_version=handle.working_dataset_version,
            duration_ms=duration_ms,
            inspection=inspection,
        )

    def _inspect_xlsx(self, path: Path) -> tuple[str, ...]:
        if not zipfile.is_zipfile(path):
            raise UnsafeFileError("XLSX upload does not have a valid ZIP signature")
        with zipfile.ZipFile(path) as archive:
            names = set(archive.namelist())
            if "[Content_Types].xml" not in names or "xl/workbook.xml" not in names:
                raise UnsafeFileError("Upload is not a structurally valid XLSX workbook")
            total_uncompressed = 0
            for member in archive.infolist():
                member_path = PurePosixPath(member.filename)
                if member_path.is_absolute() or ".." in member_path.parts:
                    raise UnsafeFileError("XLSX contains an unsafe archive path")
                total_uncompressed += member.file_size
                if total_uncompressed > self._limits.max_xlsx_uncompressed_bytes:
                    raise UnsafeFileError("XLSX uncompressed content exceeds the safety limit")
                if member.compress_size:
                    ratio = member.file_size / member.compress_size
                    if ratio > self._limits.max_xlsx_compression_ratio:
                        raise UnsafeFileError("XLSX member exceeds the compression-ratio limit")

        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
        except Exception as exc:
            raise UnsafeFileError(f"XLSX workbook cannot be opened: {exc}") from exc
        try:
            sheets = tuple(workbook.sheetnames)
        finally:
            workbook.close()
        if not sheets:
            raise UnsafeFileError("XLSX workbook contains no worksheets")
        return sheets

    def _read_frame(
        self,
        source_path: Path,
        inspection: UploadInspection,
        selected_sheet: str | None,
    ) -> pd.DataFrame:
        if inspection.file_format == "csv":
            assert inspection.encoding is not None
            assert inspection.delimiter is not None
            try:
                return pd.read_csv(
                    source_path,
                    encoding=inspection.encoding,
                    sep=inspection.delimiter,
                    dtype_backend="numpy_nullable",
                )
            except Exception as exc:
                raise UnsafeFileError(f"CSV could not be parsed: {exc}") from exc

        assert selected_sheet is not None
        workbook = None
        try:
            workbook = load_workbook(source_path, read_only=True, data_only=True)
            worksheet = workbook[selected_sheet]
            values = worksheet.iter_rows(values_only=True)
            header = next(values, None)
            if header is None:
                raise UnsafeFileError("Selected XLSX sheet is empty")
            frame = pd.DataFrame(values, columns=header)
            return frame.convert_dtypes(dtype_backend="numpy_nullable")
        except UnsafeFileError:
            raise
        except Exception as exc:
            raise UnsafeFileError(f"XLSX sheet could not be parsed: {exc}") from exc
        finally:
            if workbook is not None:
                workbook.close()

    def _write_working_database(
        self,
        frame: pd.DataFrame,
        path: Path,
        *,
        dataset_id: UUID,
        source_sha256: str,
        working_dataset_version: int,
    ) -> None:
        connection = duckdb.connect(str(path))
        try:
            connection.execute(f"SET memory_limit='{self._limits.duckdb_memory_limit_mb}MB'")
            connection.register("ingested_frame", frame)
            connection.execute("CREATE TABLE dataset AS SELECT * FROM ingested_frame")
            connection.unregister("ingested_frame")
            connection.execute(
                "CREATE TABLE _dataset_metadata ("
                "dataset_id VARCHAR NOT NULL, source_sha256 VARCHAR NOT NULL, "
                "working_dataset_version INTEGER NOT NULL)"
            )
            connection.execute(
                "INSERT INTO _dataset_metadata VALUES (?, ?, ?)",
                [str(dataset_id), source_sha256, working_dataset_version],
            )
            connection.execute("CHECKPOINT")
        finally:
            connection.close()

    def _connect(self, path: Path) -> duckdb.DuckDBPyConnection:
        if not path.is_file():
            raise DatasetIntegrityError("Working Dataset database is missing")
        connection = duckdb.connect(str(path), read_only=True)
        connection.execute(f"SET memory_limit='{self._limits.duckdb_memory_limit_mb}MB'")
        connection.execute("SET enable_external_access=false")
        connection.execute("SET lock_configuration=true")
        return connection

    def _verify_handle(self, handle: DatasetHandle) -> None:
        source_path = handle.source_path.resolve()
        working_path = handle.working_database_path.resolve()
        expected_source_directory = self._session_directory / "source"
        valid_source_identity = (
            source_path.parent == expected_source_directory
            and source_path.stem == str(handle.dataset.dataset_id)
            and source_path.suffix.casefold() in _SUPPORTED_EXTENSIONS
        )
        if not valid_source_identity:
            raise DatasetIntegrityError("Source Dataset path does not match its identity")
        expected_working_path = (
            self._session_directory
            / "working"
            / str(handle.working_dataset_version)
            / "dataset.duckdb"
        )
        if working_path != expected_working_path:
            raise DatasetIntegrityError("Working Dataset path does not match its version")
        if handle.table_name != DATASET_TABLE:
            raise DatasetIntegrityError("Working Dataset table identity is invalid")
        if not source_path.is_file():
            raise DatasetIntegrityError("Source Dataset is missing")
        if _hash_file(source_path) != handle.dataset.sha256:
            raise DatasetIntegrityError("Source Dataset hash no longer matches its identity")
        if not working_path.is_file():
            raise DatasetIntegrityError("Working Dataset is missing")
        try:
            connection = duckdb.connect(str(working_path), read_only=True)
            try:
                metadata = connection.execute(
                    "SELECT dataset_id, source_sha256, working_dataset_version "
                    "FROM _dataset_metadata"
                ).fetchall()
            finally:
                connection.close()
        except duckdb.Error as exc:
            raise DatasetIntegrityError("Working Dataset metadata is missing or invalid") from exc
        expected_metadata = [
            (
                str(handle.dataset.dataset_id),
                handle.dataset.sha256,
                handle.working_dataset_version,
            )
        ]
        if metadata != expected_metadata:
            raise DatasetIntegrityError("Working Dataset metadata does not match its handle")

    def _profile_field(
        self,
        connection: duckdb.DuckDBPyConnection,
        name: str,
        data_type: str,
        row_count: int,
    ) -> FieldProfile:
        quoted = quote_identifier(name)
        normalized_type = data_type.split("(", maxsplit=1)[0].upper()
        is_text = normalized_type in {"VARCHAR", "CHAR", "TEXT"}
        null_expression = f"{quoted} IS NULL"
        if is_text:
            null_expression = f"({quoted} IS NULL OR TRIM({quoted}) = '')"
        missing_count, unique_count = _require_row(
            connection.execute(
                f"SELECT COUNT(*) FILTER (WHERE {null_expression}), "
                f"COUNT(DISTINCT {quoted}) FILTER (WHERE NOT ({null_expression})) FROM dataset"
            ).fetchone()
        )
        missing_count = int(missing_count)
        unique_count = int(unique_count)
        non_missing = row_count - missing_count
        kind, confidence = _infer_field_kind(normalized_type, unique_count, non_missing)
        warnings: list[str] = []
        if unique_count <= 1:
            warnings.append("constant or empty field")
        elif non_missing:
            max_frequency = int(
                _require_row(
                    connection.execute(
                        f"SELECT MAX(frequency) FROM ("
                        f"SELECT COUNT(*) AS frequency FROM dataset "
                        f"WHERE NOT ({null_expression}) GROUP BY {quoted})"
                    ).fetchone()
                )[0]
            )
            if max_frequency / non_missing >= _NEAR_CONSTANT_SHARE:
                warnings.append("near-constant field (>=95% same value)")
        if (
            non_missing >= _MIN_IDENTIFIER_ROWS
            and unique_count == non_missing
            and (is_text or normalized_type in _INTEGER_TYPES)
        ):
            warnings.append(IDENTIFIER_LIKE_WARNING)
        if row_count and missing_count / row_count >= 0.5:
            warnings.append("high missingness")
        if is_text and unique_count > _MAX_CATEGORICAL_UNIQUE_COUNT:
            warnings.append("high cardinality")
        looks_like_pii = self._looks_like_pii(connection, name, quoted)
        if is_text and unique_count > 1:
            spelling_warning = _inconsistent_spelling_warning(
                connection, quoted, null_expression, include_examples=not looks_like_pii
            )
            if spelling_warning:
                warnings.append(spelling_warning)
        if looks_like_pii:
            warnings.append(POSSIBLE_PII_WARNING)

        numeric_summary = None
        temporal_summary = None
        top_values: tuple[CategoryFrequency, ...] = ()
        if kind is FieldKind.NUMERIC:
            numeric_summary = _numeric_summary(connection, quoted)
            if numeric_summary.outlier_count:
                warnings.append(
                    f"{numeric_summary.outlier_count} numeric outlier(s) detected by "
                    f"{numeric_summary.outlier_method}"
                )
        elif kind is FieldKind.DATETIME:
            earliest, latest = _require_row(
                connection.execute(f"SELECT MIN({quoted}), MAX({quoted}) FROM dataset").fetchone()
            )
            if earliest is not None and latest is not None:
                temporal_summary = TemporalSummary(
                    earliest=str(_normalize_scalar(earliest)),
                    latest=str(_normalize_scalar(latest)),
                )
        elif non_missing:
            frequency_rows = connection.execute(
                f"SELECT {quoted}, COUNT(*) AS frequency FROM dataset "
                f"WHERE NOT ({null_expression}) GROUP BY {quoted} "
                f"ORDER BY frequency DESC, CAST({quoted} AS VARCHAR) "
                f"LIMIT {self._limits.profile_top_values}"
            ).fetchall()
            top_values = tuple(
                CategoryFrequency(
                    value=_normalize_scalar(value),
                    count=int(count),
                    rate=float(count) / non_missing,
                )
                for value, count in frequency_rows
            )

        return FieldProfile(
            name=name,
            kind=kind,
            missing_count=missing_count,
            missing_rate=(missing_count / row_count) if row_count else 0.0,
            unique_count=unique_count,
            inference_confidence=confidence,
            type_ambiguous=kind is FieldKind.UNKNOWN,
            warnings=tuple(warnings),
            numeric_summary=numeric_summary,
            top_values=top_values,
            temporal_summary=temporal_summary,
        )

    def _looks_like_pii(
        self,
        connection: duckdb.DuckDBPyConnection,
        name: str,
        quoted: str,
    ) -> bool:
        normalized_name = re.sub(r"[^a-z0-9]+", "_", _ascii_fold(name)).strip("_")
        if _PII_NAME.search(normalized_name):
            return True
        samples = connection.execute(
            f"SELECT CAST({quoted} AS VARCHAR) FROM dataset WHERE {quoted} IS NOT NULL LIMIT 50"
        ).fetchall()
        texts = [str(value) for (value,) in samples]
        return any(
            _EMAIL_VALUE.match(text)
            # An ISO date such as 2026-03-02 has the digits-and-dashes shape of a phone number.
            or (_PHONE_VALUE.match(text) and not _ISO_DATE_VALUE.match(text))
            for text in texts
        )


def _ascii_fold(value: str) -> str:
    """Lowercase and strip diacritics, including the Vietnamese đ, for keyword matching."""
    decomposed = unicodedata.normalize("NFKD", value.casefold().replace("đ", "d"))
    return "".join(character for character in decomposed if not unicodedata.combining(character))


def _hash_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _make_read_only(path: Path) -> None:
    path.chmod(path.stat().st_mode & ~(stat.S_IWUSR | stat.S_IWGRP | stat.S_IWOTH))


def _unlink_read_only(path: Path) -> None:
    if not path.exists():
        return
    path.chmod(path.stat().st_mode | stat.S_IWUSR)
    path.unlink()


def _inspect_csv(path: Path) -> tuple[str, str, tuple[str, ...]]:
    with path.open("rb") as source:
        sample_bytes = source.read(128 * 1024)
    for signature, description in _BINARY_SIGNATURES.items():
        if sample_bytes.startswith(signature):
            raise UnsafeFileError(f"This .csv file is {description}, not comma-separated text")
    match = from_bytes(sample_bytes).best()
    if match is None or not match.encoding:
        raise UnsafeFileError("CSV encoding could not be detected")
    encoding = match.encoding
    try:
        sample = sample_bytes.decode(encoding)
    except UnicodeDecodeError as exc:
        raise UnsafeFileError("CSV sample cannot be decoded consistently") from exc
    if "\x00" in sample:
        raise UnsafeFileError("CSV contains NUL bytes")

    warnings: list[str] = []
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ","
        warnings.append("delimiter detection was uncertain; comma was selected")

    try:
        header = next(csv.reader(sample.splitlines(), delimiter=delimiter))
    except (csv.Error, StopIteration) as exc:
        raise UnsafeFileError("CSV header could not be read") from exc
    _validate_header_values(header)
    if encoding.casefold().replace("_", "-") not in {"utf-8", "utf-8-sig", "ascii"}:
        warnings.append(f"non-UTF-8 encoding detected: {encoding}")
    return encoding, delimiter, tuple(warnings)


def _select_sheet(sheets: tuple[str, ...], requested: str | None) -> str | None:
    if not sheets:
        if requested is not None:
            raise SheetSelectionError("CSV uploads do not accept a sheet name")
        return None
    if requested is None:
        if len(sheets) == 1:
            return sheets[0]
        raise SheetSelectionError("A sheet must be selected for a multi-sheet workbook")
    if requested not in sheets:
        raise SheetSelectionError(f"Unknown XLSX sheet: {requested}")
    return requested


def _validate_header_values(values: list[Any] | tuple[Any, ...]) -> tuple[str, ...]:
    normalized = tuple("" if value is None else str(value).strip() for value in values)
    if not normalized or any(not value for value in normalized):
        raise UnsafeFileError("Every column must have a non-empty header")
    folded = [value.casefold() for value in normalized]
    if len(folded) != len(set(folded)):
        raise UnsafeFileError("Column headers must be unique after trimming")
    return normalized


def _validate_and_normalize_headers(frame: pd.DataFrame) -> None:
    normalized = _validate_header_values(tuple(frame.columns))
    frame.columns = list(normalized)
    if frame.empty and len(frame.columns) == 0:
        raise UnsafeFileError("Dataset must contain at least one column")


def _coerce_unambiguous_dates(frame: pd.DataFrame) -> None:
    for name in frame.columns:
        series = frame[name]
        if not pd.api.types.is_string_dtype(series.dtype):
            continue
        present = series.dropna()
        if present.empty:
            continue
        # ISO dates are unambiguous in any language; other formats need a date-like column name.
        if all(_ISO_DATE_VALUE.match(str(value)) for value in present):
            date_format = "ISO8601"
        elif _DATE_NAME.search(str(name)):
            date_format = "mixed"
        else:
            continue
        parsed = pd.to_datetime(present, errors="coerce", format=date_format, utc=True)
        if parsed.notna().all():
            normalized = pd.to_datetime(series, errors="coerce", format=date_format, utc=True)
            frame[name] = normalized.dt.tz_convert(None)


def _inconsistent_spelling_warning(
    connection: duckdb.DuckDBPyConnection,
    quoted: str,
    null_expression: str,
    *,
    include_examples: bool,
) -> str | None:
    """Flag text values that differ only by letter case or surrounding spaces.

    Values are never changed; the warning lets the user and the planner see that an exact filter
    would miss rows. Examples are omitted for possible PII so no personal value is repeated.
    """
    rows = connection.execute(
        f"SELECT LOWER(TRIM({quoted})) AS spelling_key, "
        f"LIST(DISTINCT {quoted} ORDER BY {quoted}) AS spellings FROM dataset "
        f"WHERE NOT ({null_expression}) GROUP BY spelling_key "
        f"HAVING COUNT(DISTINCT {quoted}) > 1 "
        f"ORDER BY COUNT(*) DESC, spelling_key LIMIT {_MAX_SPELLING_EXAMPLES}"
    ).fetchall()
    if not rows:
        return None
    if not include_examples:
        return INCONSISTENT_SPELLING_WARNING
    examples = "; ".join(
        " / ".join(f'"{str(value)[:_MAX_SPELLING_EXAMPLE_CHARS]}"' for value in spellings[:3])
        for _, spellings in rows
    )
    return f"{INCONSISTENT_SPELLING_WARNING}: {examples}"


def _infer_field_kind(
    normalized_type: str, unique_count: int, non_missing: int
) -> tuple[FieldKind, float]:
    if normalized_type in _NUMERIC_TYPES:
        return FieldKind.NUMERIC, 1.0
    if normalized_type == "BOOLEAN":
        return FieldKind.BOOLEAN, 1.0
    if normalized_type == "DATE" or normalized_type.startswith(("TIMESTAMP", "TIME")):
        return FieldKind.DATETIME, 1.0
    if normalized_type in {"VARCHAR", "CHAR", "TEXT"}:
        ratio = unique_count / non_missing if non_missing else 0.0
        if unique_count <= _MAX_CATEGORICAL_UNIQUE_COUNT or ratio <= _MAX_CATEGORICAL_UNIQUE_RATIO:
            return FieldKind.CATEGORICAL, 0.9
        return FieldKind.SHORT_TEXT, 0.8
    return FieldKind.UNKNOWN, 0.5


def _numeric_summary(connection: duckdb.DuckDBPyConnection, quoted: str) -> NumericSummary:
    row = _require_row(
        connection.execute(
            f"SELECT COUNT({quoted}), AVG({quoted}), STDDEV_SAMP({quoted}), "
            f"MIN({quoted}), QUANTILE_CONT({quoted}, 0.25), MEDIAN({quoted}), "
            f"QUANTILE_CONT({quoted}, 0.75), MAX({quoted}) FROM dataset"
        ).fetchone()
    )
    values = [_optional_float(value) for value in row[1:]]
    first_quartile = values[3]
    third_quartile = values[5]
    outlier_count = 0
    if first_quartile is not None and third_quartile is not None:
        interquartile_range = third_quartile - first_quartile
        lower_bound = first_quartile - 1.5 * interquartile_range
        upper_bound = third_quartile + 1.5 * interquartile_range
        outlier_count = int(
            _require_row(
                connection.execute(
                    f"SELECT COUNT(*) FROM dataset WHERE {quoted} < ? OR {quoted} > ?",
                    [lower_bound, upper_bound],
                ).fetchone()
            )[0]
        )
    return NumericSummary(
        count=int(row[0]),
        mean=values[0],
        standard_deviation=values[1],
        minimum=values[2],
        first_quartile=values[3],
        median=values[4],
        third_quartile=values[5],
        maximum=values[6],
        outlier_count=outlier_count,
    )


def _optional_float(value: Any) -> float | None:
    if value is None:
        return None
    result = float(value)
    return result if math.isfinite(result) else None


def _normalize_scalar(value: Any) -> str | int | float | bool | None:
    if value is None or value is pd.NA:
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, int):
        return value
    if isinstance(value, (float, Decimal)):
        number = float(value)
        return number if math.isfinite(number) else None
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value)


# Uniqueness says little about a tiny table, and continuous measures are naturally unique.
_MIN_IDENTIFIER_ROWS = 20
# A text field with at most this many distinct values, or distinct values in at most this share
# of its rows, is categorical; beyond the count it also earns a high-cardinality warning.
_MAX_CATEGORICAL_UNIQUE_COUNT = 50
_MAX_CATEGORICAL_UNIQUE_RATIO = 0.2
# A field whose most frequent value covers at least this share of rows is near-constant.
_NEAR_CONSTANT_SHARE = 0.95
_INTEGER_TYPES = frozenset(
    {
        "TINYINT",
        "SMALLINT",
        "INTEGER",
        "BIGINT",
        "HUGEINT",
        "UTINYINT",
        "USMALLINT",
        "UINTEGER",
        "UBIGINT",
    }
)


def _require_row(row: tuple[Any, ...] | None) -> tuple[Any, ...]:
    if row is None:
        raise DatasetIntegrityError("Internal analytical query returned no result row")
    return row


def _normalize_rows_with_deadline(
    raw_rows: list[tuple[Any, ...]], *, deadline: float
) -> tuple[tuple[Any, ...], ...]:
    """Normalize database values while allowing the caller's full-read deadline to expire."""
    rows: list[tuple[Any, ...]] = []
    for index, row in enumerate(raw_rows):
        if index % 256 == 0 and time.perf_counter() >= deadline:
            raise QueryTimeoutError("Full-data read exceeded its end-to-end deadline")
        rows.append(tuple(_normalize_scalar(value) for value in row))
    if time.perf_counter() >= deadline:
        raise QueryTimeoutError("Full-data read exceeded its end-to-end deadline")
    return tuple(rows)
